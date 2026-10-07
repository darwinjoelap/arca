"""Pruebas de la exportación y la carga desde Excel (solo el director)."""

from datetime import date
from decimal import Decimal
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from core.models import RegistroAuditoria
from finanzas.models import Caja, Concepto, Cuenta, Movimiento
from finanzas.services import saldo_cuenta
from finanzas.tests import LibroMixin
from inventario import services as inventario
from inventario.models import Articulo

from . import datos
from .models import Membresia, TipoMiembro

D = Decimal


def libro(**hojas):
    """Un .xlsx con el encabezado de la plantilla y las filas dadas por hoja."""
    wb = Workbook()
    for titulo, filas in hojas.items():
        hoja = wb.create_sheet(titulo)
        hoja.append(["Nota de ayuda"])
        hoja.append([c for c, _ in datos.HOJAS[titulo][0]])
        for fila in filas:
            hoja.append(list(fila))
    del wb[wb.sheetnames[0]]
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


class ImportarTests(LibroMixin, TestCase):
    def setUp(self):
        self.m_director = Membresia.objects.get(organizacion=self.org, es_dueno=True)

    def cargar(self, contenido, simular=False):
        return datos.importar(self.org, contenido, membresia=self.m_director, simular=simular)

    def completo(self):
        return libro(
            Miembros=[("Ana Gil", "Vecino")],
            Cuentas=[("Fondo obra", "Zelle", "banco", "usd", "BofA", 120)],
            Conceptos=[("egreso", "Luz", "Servicios")],
            Movimientos=[
                (self.hoy, "ingreso", "Zelle", "1.250,50", "Donaciones", "", "Ana Gil", "", ""),
                (self.hoy.strftime("%d/%m/%Y"), "Egreso", "Fondo obra / Zelle", 30, "Luz", "Recibo", "", "Corpoelec", "77"),
            ],
            Inventario=[("consumible", "Cloro", "", "Limpieza", "litro", 5, 2, "USD", "Depósito", 12),
                        ("bien", "Silla", "S-1", "", "", "", "", "", "", "")],
        )

    def test_la_simulacion_cuenta_pero_no_guarda(self):
        r = self.cargar(self.completo(), simular=True)
        self.assertEqual(r.errores, [])
        self.assertFalse(r.guardado)
        self.assertEqual(r.creados["Movimientos"], 2)
        self.assertFalse(Cuenta.objects.filter(nombre="Zelle").exists())
        self.assertFalse(Membresia.objects.filter(nombre_visible="Ana Gil").exists())

    def test_carga_completa(self):
        r = self.cargar(self.completo())
        self.assertTrue(r.guardado, r.errores)
        zelle = Cuenta.objects.get(organizacion=self.org, nombre="Zelle")
        self.assertEqual((zelle.caja.nombre, zelle.moneda, zelle.tipo), ("Fondo obra", "USD", "banco"))
        self.assertEqual(saldo_cuenta(zelle), D("120") + D("1250.50") - D("30"))
        ana = Membresia.objects.get(organizacion=self.org, nombre_visible="Ana Gil")
        self.assertIsNone(ana.user_id)                      # entra sin acceso
        self.assertEqual(ana.tipo.nombre, "Vecino")
        self.assertFalse(ana.tipo.puede_registrar_egresos)  # tipo nuevo, sin permisos
        luz = Concepto.objects.get(organizacion=self.org, nombre="Luz")
        self.assertEqual(luz.padre.nombre, "Servicios")
        ingreso = Movimiento.objects.get(organizacion=self.org, cuenta=zelle, tipo="ingreso")
        self.assertEqual((ingreso.miembro, ingreso.concepto.nombre, ingreso.estado), (ana, "Donaciones", "confirmado"))
        cloro = Articulo.objects.get(organizacion=self.org, nombre="Cloro")
        self.assertEqual(inventario.totales_por_articulo(self.org)[cloro.pk], D("12"))
        self.assertEqual(r.creados["Cajas"], 1)
        self.assertEqual(r.creados["Conceptos"], 3)   # Servicios, Luz y Donaciones

    def test_repetir_el_archivo_no_duplica(self):
        self.cargar(self.completo())
        antes = Movimiento.objects.filter(organizacion=self.org).count()
        r = self.cargar(self.completo())
        self.assertEqual(r.total, 0)
        self.assertEqual(r.omitidos["Movimientos"], 2)
        self.assertEqual(Movimiento.objects.filter(organizacion=self.org).count(), antes)

    def test_dos_filas_iguales_en_el_mismo_archivo_si_son_dos_movimientos(self):
        fila = (self.hoy, "ingreso", "Efectivo", 5, "Aportes", "", "", "", "")
        r = self.cargar(libro(Movimientos=[fila, fila]))
        self.assertEqual(r.creados["Movimientos"], 2)

    def test_un_error_impide_toda_la_carga(self):
        contenido = libro(
            Cuentas=[("Nueva", "Caja fuerte", "efectivo", "USD", "", 0)],
            Movimientos=[(self.hoy, "ingreso", "Efectivo", 5, "Aportes", "", "", "", ""),
                         (self.hoy, "ingreso", "No existe", 5, "Aportes", "", "", "", ""),
                         ("ayer", "egreso", "Efectivo", "mucho", "", "", "", "", ""),
                         (date(2001, 1, 1), "egreso", "Efectivo", 5, "Comida", "", "", "", "")],
        )
        r = self.cargar(contenido)
        self.assertFalse(r.guardado)
        self.assertEqual([(h, f) for h, f, _ in r.errores], [("Movimientos", 4), ("Movimientos", 5), ("Movimientos", 6)])
        self.assertIn("No existe la cuenta", r.errores[0][2])
        self.assertIn("ejercicio", r.errores[2][2])
        self.assertFalse(Caja.objects.filter(nombre="Nueva").exists())
        self.assertFalse(Movimiento.objects.filter(organizacion=self.org).exists())

    def test_no_toca_a_la_otra_organizacion(self):
        # «Caja chica» es de B: desde A no se encuentra.
        r = self.cargar(libro(Movimientos=[(self.hoy, "ingreso", "Caja chica", 5, "", "Algo", "", "", "")]))
        self.assertEqual(len(r.errores), 1)
        self.assertFalse(Movimiento.objects.filter(organizacion=self.org_b).exists())
        r = self.cargar(libro(Conceptos=[("egreso", "Canchas", "")]))
        self.assertTrue(r.guardado)   # mismo nombre que en B: se crea aparte, en A
        self.assertEqual(Concepto.objects.filter(nombre="Canchas").count(), 2)

    def test_archivo_que_no_es_la_plantilla(self):
        with self.assertRaises(datos.ArchivoInvalido):
            self.cargar(b"esto no es un excel")
        wb = Workbook()
        salida = BytesIO()
        wb.save(salida)
        with self.assertRaises(datos.ArchivoInvalido):
            self.cargar(salida.getvalue())

    def test_la_plantilla_con_sus_ejemplos_se_puede_cargar(self):
        from cambio.models import TasaCambio
        TasaCambio.objects.create(fecha=date(2026, 1, 1), valor=D("90"), fuente="bcv")
        from .models import Ejercicio
        if not Ejercicio.objects.filter(organizacion=self.org, fecha_inicio__lte=date(2026, 1, 15), fecha_fin__gte=date(2026, 1, 20)).exists():
            Ejercicio.objects.create(organizacion=self.org, nombre="2026", fecha_inicio=date(2026, 1, 1), fecha_fin=date(2026, 12, 31))
        r = self.cargar(datos.plantilla(), simular=True)
        self.assertEqual(r.errores, [])
        self.assertEqual(r.creados["Movimientos"], 2)


class VistasDatosTests(LibroMixin, TestCase):
    def test_solo_el_director(self):
        for nombre in ("datos", "datos_exportar", "datos_plantilla"):
            self.entrar("tesorero")
            self.assertEqual(self.client.get(reverse(f"organizaciones:{nombre}")).status_code, 403)
        admin = Membresia.objects.get(organizacion=self.org, user__username="tesorero")
        admin.es_administrador = True
        admin.save()
        self.assertEqual(self.client.get(reverse("organizaciones:datos_exportar")).status_code, 403)
        self.entrar("dir_a")
        self.assertEqual(self.client.get(reverse("organizaciones:datos")).status_code, 200)

    def test_exportar_trae_solo_lo_propio(self):
        self.mov("ingreso", "25")
        self.mov("egreso", "7", org=self.org_b, cuenta=self.cuenta_b, concepto=self.concepto_b,
                 membresia=Membresia.objects.get(organizacion=self.org_b, es_dueno=True))
        self.entrar("dir_a")
        respuesta = self.client.get(reverse("organizaciones:datos_exportar"))
        self.assertEqual(respuesta.status_code, 200)
        wb = load_workbook(BytesIO(respuesta.content))
        self.assertEqual(wb.sheetnames, ["Organización", "Miembros", "Cuentas", "Conceptos", "Ejercicios", "Movimientos",
                                         "Traslados", "Presupuesto", "Cuotas", "Anticipos", "Arqueos", "Inventario",
                                         "Mov. de inventario"])
        todo = " ".join(str(c) for hoja in wb for fila in hoja.iter_rows(values_only=True) for c in fila if c is not None)
        self.assertIn("Aportes", todo)
        for ajeno in ("Club B", "Caja chica", "Canchas"):
            self.assertNotIn(ajeno, todo)
        filas = list(wb["Cuentas"].iter_rows(values_only=True))
        efectivo = next(f for f in filas if f[1] == "Efectivo")
        self.assertEqual((efectivo[5], efectivo[6]), (50, 75))    # saldo inicial y actual
        self.assertTrue(RegistroAuditoria.objects.filter(organizacion=self.org, accion="exportar_datos").exists())

    def test_revisar_y_confirmar(self):
        self.entrar("dir_a")
        archivo = SimpleUploadedFile("mio.xlsx", libro(Miembros=[("Luis Mora", "Vecino")]))
        respuesta = self.client.post(reverse("organizaciones:datos_cargar"), {"archivo": archivo})
        self.assertContains(respuesta, "Confirmar y cargar")
        self.assertFalse(Membresia.objects.filter(nombre_visible="Luis Mora").exists())
        respuesta = self.client.post(reverse("organizaciones:datos_cargar"), {"confirmar": "1"})
        self.assertRedirects(respuesta, reverse("organizaciones:datos"))
        self.assertTrue(Membresia.objects.filter(organizacion=self.org, nombre_visible="Luis Mora").exists())
        self.assertTrue(TipoMiembro.objects.filter(organizacion=self.org, nombre="Vecino").exists())
        self.assertTrue(RegistroAuditoria.objects.filter(organizacion=self.org, accion="importar_datos").exists())
        # Confirmar dos veces no repite: el archivo ya salió de la sesión.
        respuesta = self.client.post(reverse("organizaciones:datos_cargar"), {"confirmar": "1"}, follow=True)
        self.assertContains(respuesta, "La revisión venció")

    def test_con_errores_no_ofrece_confirmar(self):
        self.entrar("dir_a")
        archivo = SimpleUploadedFile("mio.xlsx", libro(Cuentas=[("Casa", "Otra", "efectivo", "euros", "", 0)]))
        respuesta = self.client.post(reverse("organizaciones:datos_cargar"), {"archivo": archivo})
        self.assertContains(respuesta, "La moneda debe ser")
        self.assertNotContains(respuesta, "Confirmar y cargar")
