"""Entrada y salida de datos de UNA organización en Excel.

- `exportar(organizacion)`: un libro con una hoja por tema (todo lo de la
  organización, menos las cuentas personales de los miembros, que son privadas).
- `plantilla()`: el libro vacío, con ejemplos, que se llena para migrar.
- `importar(organizacion, contenido, membresia, simular)`: lee ese libro.
  Siempre corre completo dentro de una transacción; si es simulación o hay
  algún error, se deshace todo. Es todo o nada.

Repetir un archivo no duplica: lo que ya existe (por nombre) se omite, y cada
movimiento lleva una huella (`uuid_cliente`) derivada de su contenido.
"""

import uuid
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from finanzas.models import Anticipo, Arqueo, Caja, Concepto, Cuenta, CuotaMiembro, Movimiento, Traslado
from finanzas.services import _sumas_por_cuenta, registrar_movimiento
from inventario import services as inventario
from inventario.models import Articulo, CategoriaArticulo, MovimientoInventario, Ubicacion
from presupuestos.models import PartidaPresupuestaria

from .models import Ejercicio, Membresia, TipoMiembro

CARBON, MARFIL, PIEDRA = "1C1B1A", "F6F3EC", "6F6A63"
MAX_FILAS = 5000          # por hoja, al importar
MAX_BYTES = 2 * 1024 * 1024
ESPACIO = uuid.UUID("8f0c1c0e-5a37-4a55-9d0e-6a7263612d31")  # para las huellas de los movimientos


# --- Hojas -------------------------------------------------------------------


def _hoja(libro, titulo, columnas, filas, nota=""):
    """Agrega una hoja con encabezado fijo. `columnas` = [(título, ancho)]."""
    hoja = libro.create_sheet(titulo[:31])
    inicio = 1
    if nota:
        hoja["A1"] = nota
        hoja["A1"].font = Font(italic=True, color=PIEDRA, size=10)
        hoja["A1"].alignment = Alignment(wrap_text=True, vertical="top")
        hoja.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(columnas), 2))
        hoja.row_dimensions[1].height = 48
        inicio = 2
    for i, (nombre, ancho) in enumerate(columnas, start=1):
        celda = hoja.cell(row=inicio, column=i, value=nombre)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=CARBON)
        celda.alignment = Alignment(vertical="center", wrap_text=True)
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    for fila in filas:
        hoja.append([_celda(v) for v in fila])
    hoja.freeze_panes = hoja.cell(row=inicio + 1, column=1)
    return hoja


def _celda(valor):
    if isinstance(valor, Decimal):
        return float(valor)
    if isinstance(valor, datetime):
        return timezone.localtime(valor).replace(tzinfo=None) if timezone.is_aware(valor) else valor
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    return valor


def _bytes(libro):
    del libro[libro.sheetnames[0]]  # la hoja vacía con la que nace el libro
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


# --- Exportación -------------------------------------------------------------


def exportar(organizacion):
    """El libro completo de la organización, como bytes de .xlsx."""
    o = organizacion
    libro = Workbook()

    _hoja(libro, "Organización", [("Dato", 28), ("Valor", 60)], [
        ("Nombre", o.nombre), ("Enlace", f"/{o.slug}/"), ("RIF", o.rif), ("Dirección", o.direccion),
        ("Teléfono", o.telefono), ("Correo", o.email), ("Moneda base", o.moneda_base),
        ("Los egresos requieren aprobación", o.requiere_aprobacion_egresos),
        ("Exportado el", timezone.localtime().strftime("%d/%m/%Y %H:%M")),
    ])

    miembros = Membresia.objects.filter(organizacion=o).select_related("user", "tipo")
    _hoja(libro, "Miembros", [("Nombre", 34), ("Rol o tipo", 24), ("Usuario", 20), ("Activo", 10), ("Desde", 14)], [
        (m.nombre, m.rol_visible, m.usuario_texto, m.activa, m.creada_en.date()) for m in miembros
    ])

    sumas = _sumas_por_cuenta(o)
    cuentas = Cuenta.objects.filter(organizacion=o).select_related("caja").order_by("caja__nombre", "nombre")
    _hoja(libro, "Cuentas", [("Caja", 24), ("Cuenta", 28), ("Tipo", 12), ("Moneda", 10), ("Banco", 20),
                              ("Saldo inicial", 16), ("Saldo actual", 16), ("Activa", 10)], [
        (c.caja.nombre, c.nombre, c.get_tipo_display(), c.moneda, c.banco, c.saldo_inicial,
         c.saldo_inicial + sumas.get(c.pk, Decimal("0")), c.activa) for c in cuentas
    ])

    conceptos = Concepto.objects.filter(organizacion=o).select_related("padre", "caja")
    _hoja(libro, "Conceptos", [("Tipo", 12), ("Nombre", 32), ("Agrupado en", 28), ("Caja", 22), ("Activo", 10)], [
        (c.get_tipo_display(), c.nombre, c.padre.nombre if c.padre_id else "", c.caja.nombre if c.caja_id else "", c.activo)
        for c in conceptos
    ])

    ejercicios = Ejercicio.objects.filter(organizacion=o).order_by("fecha_inicio")
    _hoja(libro, "Ejercicios", [("Nombre", 16), ("Desde", 14), ("Hasta", 14), ("Activo", 10), ("Cerrado", 10)], [
        (e.nombre, e.fecha_inicio, e.fecha_fin, e.activo, e.cerrado) for e in ejercicios
    ])

    movimientos = (
        Movimiento.objects.filter(organizacion=o)
        .select_related("caja", "cuenta", "concepto", "concepto__padre", "miembro", "miembro__user", "ejercicio",
                        "registrado_por")
        .order_by("fecha", "pk")
    )
    _hoja(libro, "Movimientos", [
        ("Fecha", 12), ("Tipo", 10), ("Vale n.º", 9), ("Caja", 20), ("Cuenta", 22), ("Concepto", 30),
        ("Descripción", 36), ("Miembro", 26), ("Tercero", 24), ("Referencia", 18), ("Moneda", 9), ("Monto", 15),
        ("Tasa", 12), ("En Bs.", 16), ("En USD", 14), ("Estado", 14), ("Ejercicio", 12), ("Registrado por", 18),
        ("Motivo de anulación", 30),
    ], [
        (m.fecha, m.get_tipo_display(), m.numero_vale, m.caja.nombre, m.cuenta.nombre,
         str(m.concepto) if m.concepto_id else "", m.descripcion, m.miembro.nombre if m.miembro_id else "", m.tercero,
         m.referencia, m.moneda, m.monto, m.tasa_aplicada, m.monto_ves, m.monto_usd, m.get_estado_display(),
         m.ejercicio.nombre, str(m.registrado_por), m.motivo_anulacion)
        for m in movimientos.iterator(chunk_size=2000)
    ])

    traslados = Traslado.objects.filter(organizacion=o).select_related(
        "cuenta_origen", "cuenta_origen__caja", "cuenta_destino", "cuenta_destino__caja").order_by("fecha", "pk")
    _hoja(libro, "Traslados", [("Fecha", 12), ("Sale de", 34), ("Moneda", 9), ("Monto que sale", 16),
                                ("Entra a", 34), ("Moneda ", 9), ("Monto que entra", 16), ("Descripción", 34),
                                ("Referencia", 18), ("Estado", 12)], [
        (t.fecha, f"{t.cuenta_origen.caja.nombre} / {t.cuenta_origen.nombre}", t.cuenta_origen.moneda, t.monto_origen,
         f"{t.cuenta_destino.caja.nombre} / {t.cuenta_destino.nombre}", t.cuenta_destino.moneda, t.monto_destino,
         t.descripcion, t.referencia, t.get_estado_display()) for t in traslados
    ])

    partidas = PartidaPresupuestaria.objects.filter(organizacion=o).select_related(
        "presupuesto", "presupuesto__caja", "presupuesto__ejercicio", "concepto"
    ).order_by("presupuesto__ejercicio__fecha_inicio", "presupuesto__caja__nombre", "concepto__nombre")
    _hoja(libro, "Presupuesto", [("Ejercicio", 12), ("Caja", 22), ("Estado", 12), ("Moneda", 9), ("Tipo", 10),
                                  ("Concepto", 30), ("Monto mensual", 16), ("Requiere solicitud", 12), ("Nota", 30)], [
        (p.presupuesto.ejercicio.nombre, p.presupuesto.caja.nombre, p.presupuesto.get_estado_display(),
         p.presupuesto.moneda, p.concepto.get_tipo_display(), p.concepto.nombre, p.monto_mensual,
         p.requiere_solicitud, p.nota) for p in partidas
    ])

    cuotas = CuotaMiembro.objects.filter(organizacion=o).select_related("membresia", "membresia__user", "concepto")
    _hoja(libro, "Cuotas", [("Miembro", 30), ("Concepto", 26), ("Moneda", 9), ("Monto mensual", 16),
                             ("Desde", 12), ("Hasta", 12), ("Nota", 30)], [
        (c.membresia.nombre, c.concepto.nombre, c.moneda, c.monto, c.vigente_desde, c.vigente_hasta, c.nota)
        for c in cuotas
    ])

    anticipos = Anticipo.objects.filter(organizacion=o).select_related(
        "cuenta", "cuenta__caja", "miembro", "miembro__user").order_by("fecha", "pk")
    _hoja(libro, "Anticipos", [("Fecha", 12), ("Cuenta", 32), ("Moneda", 9), ("Monto", 15), ("Responsable", 28),
                                ("Para qué", 36), ("Estado", 12), ("Rendido el", 12), ("Gastado", 15)], [
        (a.fecha, f"{a.cuenta.caja.nombre} / {a.cuenta.nombre}", a.cuenta.moneda, a.monto, a.responsable, a.motivo,
         a.get_estado_display(), a.fecha_rendicion, a.gastado) for a in anticipos
    ])

    arqueos = Arqueo.objects.filter(organizacion=o).select_related("cuenta", "cuenta__caja", "realizado_por").order_by("fecha", "pk")
    _hoja(libro, "Arqueos", [("Fecha", 12), ("Cuenta", 32), ("Moneda", 9), ("Según el sistema", 17), ("Contado", 15),
                              ("Diferencia", 15), ("Ajustado", 10), ("Nota", 30), ("Hecho por", 18)], [
        (a.fecha, f"{a.cuenta.caja.nombre} / {a.cuenta.nombre}", a.cuenta.moneda, a.saldo_sistema, a.contado,
         a.diferencia, a.ajuste_id is not None, a.nota, str(a.realizado_por)) for a in arqueos
    ])

    articulos, _ = inventario.resumen(o)
    _hoja(libro, "Inventario", [("Clase", 12), ("Artículo", 34), ("Código", 16), ("Categoría", 22), ("Unidad", 10),
                                 ("Existencia", 12), ("Mínimo", 10), ("Estado", 10), ("Responsable", 26),
                                 ("Valor unitario", 14), ("Moneda", 9), ("Activo", 9), ("Nota", 30)], [
        (a.get_clase_display(), a.nombre, a.codigo, a.categoria.nombre if a.categoria_id else "", a.unidad,
         a.existencia, a.minimo, "" if a.es_consumible else a.get_estado_display(),
         a.responsable.nombre if a.responsable_id else "", a.valor_unitario, a.moneda, a.activo, a.nota)
        for a in sorted(articulos, key=lambda a: a.nombre.lower())
    ])

    mov_inv = MovimientoInventario.objects.filter(organizacion=o).select_related(
        "articulo", "origen", "destino", "registrado_por").order_by("fecha", "pk")
    _hoja(libro, "Mov. de inventario", [("Fecha", 12), ("Tipo", 10), ("Motivo", 26), ("Artículo", 34),
                                         ("Cantidad", 12), ("Sale de", 22), ("Entra a", 22), ("Nota", 30),
                                         ("Registrado por", 18), ("Anulado", 9)], [
        (m.fecha, m.get_tipo_display(), m.get_motivo_display(), m.articulo.nombre, m.cantidad,
         m.origen.nombre if m.origen_id else "", m.destino.nombre if m.destino_id else "", m.nota,
         str(m.registrado_por), m.anulado) for m in mov_inv.iterator(chunk_size=2000)
    ])
    return _bytes(libro)


# --- Plantilla de migración ----------------------------------------------------

HOJAS = {
    "Miembros": (
        [("Nombre", 34), ("Tipo de miembro", 26)],
        "Personas de la comunidad. Entran SIN acceso al sistema (se les da después, una por una). "
        "Si el tipo no existe, se crea sin permisos.",
        [("María Pérez", "Residente"), ("José Rondón", "Residente")],
    ),
    "Cuentas": (
        [("Caja", 24), ("Cuenta", 28), ("Tipo", 12), ("Moneda", 10), ("Banco", 20), ("Saldo inicial", 16)],
        "Dónde está el dinero. Tipo: efectivo o banco. Moneda: USD o VES. La caja se crea si no existe. "
        "Saldo inicial: lo que había antes del primer movimiento que vas a cargar.",
        [("Caja general", "Efectivo USD", "efectivo", "USD", "", 250), ("Caja general", "Banesco", "banco", "VES", "Banesco", 0)],
    ),
    "Conceptos": (
        [("Tipo", 12), ("Nombre", 32), ("Agrupado en", 28)],
        "La lista de ingresos y egresos. Tipo: ingreso o egreso. «Agrupado en» es opcional (ej.: Luz dentro de Servicios).",
        [("ingreso", "Cuota mensual", ""), ("egreso", "Servicios", ""), ("egreso", "Luz", "Servicios")],
    ),
    "Movimientos": (
        [("Fecha", 12), ("Tipo", 10), ("Cuenta", 28), ("Monto", 14), ("Concepto", 28), ("Descripción", 36),
         ("Miembro", 26), ("Tercero", 24), ("Referencia", 18)],
        "Ingresos y egresos, uno por fila. Fecha: día/mes/año. El monto va en la moneda de la cuenta. Si dos cajas "
        "tienen una cuenta con el mismo nombre, escribe «Caja / Cuenta». El concepto se crea si no existe; sin "
        "concepto, la descripción es obligatoria.",
        [(date(2026, 1, 15), "ingreso", "Efectivo USD", 20, "Cuota mensual", "", "María Pérez", "", ""),
         (date(2026, 1, 20), "egreso", "Banesco", 1500, "Luz", "Recibo de enero", "", "Corpoelec", "000123")],
    ),
    "Inventario": (
        [("Clase", 12), ("Artículo", 34), ("Código", 16), ("Categoría", 22), ("Unidad", 10), ("Mínimo", 10),
         ("Valor unitario", 14), ("Moneda", 9), ("Ubicación", 22), ("Cantidad", 10)],
        "Bienes y consumibles. Clase: bien o consumible. El mínimo solo aplica a consumibles. Si pones cantidad, "
        "la ubicación es obligatoria (se crea si no existe) y entra como existencia inicial.",
        [("bien", "Silla plástica", "", "Mobiliario", "unidad", "", 12, "USD", "Salón", 40),
         ("consumible", "Cloro", "", "Limpieza", "litro", 5, 2, "USD", "Depósito", 12)],
    ),
}


def plantilla():
    libro = Workbook()
    for titulo, (columnas, nota, ejemplos) in HOJAS.items():
        _hoja(libro, titulo, columnas, ejemplos, nota=nota)
    return _bytes(libro)


# --- Importación -------------------------------------------------------------


class ArchivoInvalido(Exception):
    """El archivo no se puede leer como plantilla de Arca."""


@dataclass
class Resultado:
    creados: dict = field(default_factory=dict)     # {"Miembros": 3, ...}
    omitidos: dict = field(default_factory=dict)    # ya existían
    errores: list = field(default_factory=list)     # [(hoja, fila, mensaje)]
    avisos: list = field(default_factory=list)
    guardado: bool = False

    @property
    def total(self):
        return sum(self.creados.values())

    def resumen(self):
        nombres = list(dict.fromkeys(list(self.creados) + list(self.omitidos)))
        return [(n, self.creados.get(n, 0), self.omitidos.get(n, 0)) for n in nombres]


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return " ".join(str(valor).split())


def _clave(valor):
    return _texto(valor).casefold()


def _decimal(valor, campo, obligatorio=True):
    if valor is None or _texto(valor) == "":
        if obligatorio:
            raise ValueError(f"Falta {campo}.")
        return None
    if isinstance(valor, bool):
        raise ValueError(f"{campo} no es un número.")
    if isinstance(valor, (int, float)):
        return Decimal(str(valor))
    texto = _texto(valor).replace(" ", "")
    if "," in texto:  # formato venezolano: 1.234,56
        texto = texto.replace(".", "").replace(",", ".")
    try:
        return Decimal(texto)
    except InvalidOperation:
        raise ValueError(f"{campo} no es un número: «{valor}».")


def _fecha(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = _texto(valor)
    if not texto:
        raise ValueError("Falta la fecha.")
    for formato in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    raise ValueError(f"No se entiende la fecha «{texto}». Usa día/mes/año.")


def _opcion(valor, opciones, campo):
    """`opciones` = {texto aceptado: valor interno}."""
    clave = _clave(valor)
    if clave not in opciones:
        raise ValueError(f"{campo} debe ser {' o '.join(sorted(set(opciones.values())))}: «{_texto(valor)}».")
    return opciones[clave]


def _mensaje(error):
    if isinstance(error, ValidationError):
        if hasattr(error, "message_dict"):
            return " ".join(m for mensajes in error.message_dict.values() for m in mensajes)
        return " ".join(error.messages)
    return str(error)


def _filas(libro, titulo):
    """[(n.º de fila en Excel, {columna: valor})] de una hoja, o [] si no está.
    El encabezado es la primera fila cuya primera celda es la primera columna esperada."""
    if titulo not in libro.sheetnames:
        return []
    columnas = [c for c, _ in HOJAS[titulo][0]]
    hoja, encabezado, filas = libro[titulo], None, []
    for numero, valores in enumerate(hoja.iter_rows(values_only=True), start=1):
        valores = list(valores)
        if encabezado is None:
            if valores and _clave(valores[0]) == _clave(columnas[0]):
                encabezado = {_clave(v): i for i, v in enumerate(valores) if v is not None}
                faltan = [c for c in columnas if _clave(c) not in encabezado]
                if faltan:
                    raise ArchivoInvalido(f"En la hoja «{titulo}» faltan las columnas: {', '.join(faltan)}.")
            elif numero > 5:
                raise ArchivoInvalido(f"La hoja «{titulo}» no tiene el encabezado de la plantilla.")
            continue
        if all(_texto(v) == "" for v in valores):
            continue
        if len(filas) >= MAX_FILAS:
            raise ArchivoInvalido(f"La hoja «{titulo}» tiene más de {MAX_FILAS} filas: divídela en varios archivos.")
        filas.append((numero, {c: (valores[encabezado[_clave(c)]] if encabezado[_clave(c)] < len(valores) else None)
                               for c in columnas}))
    return filas


TIPOS_CUENTA = {"efectivo": "efectivo", "banco": "banco"}
MONEDAS = {"usd": "USD", "$": "USD", "dólares": "USD", "dolares": "USD",
           "ves": "VES", "bs": "VES", "bs.": "VES", "bolívares": "VES", "bolivares": "VES"}
TIPOS_MOV = {"ingreso": "ingreso", "egreso": "egreso"}
CLASES = {"bien": "bien", "consumible": "consumible"}


class _Importador:
    def __init__(self, organizacion, membresia):
        self.o, self.membresia, self.r = organizacion, membresia, Resultado()
        self.hoy = timezone.localdate()

    # Cada fila corre en su propio punto de guardado: un error no arrastra a las demás.
    def _cada(self, libro, titulo, funcion):
        for numero, fila in _filas(libro, titulo):
            try:
                with transaction.atomic():
                    creado = funcion(fila)
            except (ValueError, ValidationError) as e:
                self.r.errores.append((titulo, numero, _mensaje(e)))
                continue
            except IntegrityError:
                self.r.errores.append((titulo, numero, "Choca con algo que ya existe con ese mismo nombre."))
                continue
            destino = self.r.creados if creado else self.r.omitidos
            destino[titulo] = destino.get(titulo, 0) + 1

    def _sumar(self, nombre):
        self.r.creados[nombre] = self.r.creados.get(nombre, 0) + 1

    # -- Miembros
    def miembro(self, fila):
        nombre = _texto(fila["Nombre"])
        if not nombre:
            raise ValueError("Falta el nombre.")
        if self._buscar_miembro(nombre) is not None:
            return False
        tipo_nombre = _texto(fila["Tipo de miembro"])
        if not tipo_nombre:
            raise ValueError("Falta el tipo de miembro.")
        tipo = TipoMiembro.objects.filter(organizacion=self.o, nombre__iexact=tipo_nombre).first()
        if tipo is None:
            tipo = TipoMiembro.objects.create(organizacion=self.o, nombre=tipo_nombre)
            self._sumar("Tipos de miembro")
        m = Membresia(organizacion=self.o, tipo=tipo, nombre_visible=nombre[:120])
        m.full_clean()
        m.save()
        return True

    def _buscar_miembro(self, nombre):
        clave = _clave(nombre)
        for m in Membresia.objects.filter(organizacion=self.o).select_related("user"):
            if _clave(m.nombre) == clave or (m.user_id and m.user.username.casefold() == clave):
                return m
        return None

    # -- Cuentas
    def cuenta(self, fila):
        caja_nombre, nombre = _texto(fila["Caja"]), _texto(fila["Cuenta"])
        if not caja_nombre or not nombre:
            raise ValueError("Faltan la caja o el nombre de la cuenta.")
        caja = Caja.objects.filter(organizacion=self.o, nombre__iexact=caja_nombre).first()
        if caja is None:
            caja = Caja.objects.create(organizacion=self.o, nombre=caja_nombre[:80])
            self._sumar("Cajas")
        if Cuenta.objects.filter(organizacion=self.o, caja=caja, nombre__iexact=nombre).exists():
            return False
        cuenta = Cuenta(
            organizacion=self.o, caja=caja, nombre=nombre,
            tipo=_opcion(fila["Tipo"] or "efectivo", TIPOS_CUENTA, "El tipo"),
            moneda=_opcion(fila["Moneda"], MONEDAS, "La moneda"), banco=_texto(fila["Banco"]),
            saldo_inicial=_decimal(fila["Saldo inicial"], "el saldo inicial", obligatorio=False) or Decimal("0"),
        )
        cuenta.full_clean()
        cuenta.save()
        return True

    def _buscar_cuenta(self, texto):
        texto = _texto(texto)
        if not texto:
            raise ValueError("Falta la cuenta.")
        cuentas = Cuenta.objects.filter(organizacion=self.o).select_related("caja")
        if "/" in texto:
            caja, _, nombre = (p.strip() for p in texto.partition("/"))
            hallada = [c for c in cuentas if _clave(c.caja.nombre) == _clave(caja) and _clave(c.nombre) == _clave(nombre)]
        else:
            hallada = [c for c in cuentas if _clave(c.nombre) == _clave(texto)]
        if not hallada:
            raise ValueError(f"No existe la cuenta «{texto}». Agrégala en la hoja Cuentas.")
        if len(hallada) > 1:
            raise ValueError(f"Hay más de una cuenta «{texto}»: escribe «Caja / Cuenta».")
        return hallada[0]

    # -- Conceptos
    def _concepto(self, tipo, nombre, padre_nombre=""):
        """(concepto, creado). Crea el grupo si hace falta."""
        existente = Concepto.objects.filter(organizacion=self.o, tipo=tipo, nombre__iexact=nombre).first()
        if existente is not None:
            return existente, False
        padre = None
        if padre_nombre:
            padre, creado = self._concepto(tipo, padre_nombre)
            if creado:
                self._sumar("Conceptos")
        concepto = Concepto(organizacion=self.o, tipo=tipo, nombre=nombre, padre=padre)
        concepto.full_clean()
        concepto.save()
        return concepto, True

    def concepto(self, fila):
        nombre = _texto(fila["Nombre"])
        if not nombre:
            raise ValueError("Falta el nombre.")
        _, creado = self._concepto(_opcion(fila["Tipo"], TIPOS_MOV, "El tipo"), nombre, _texto(fila["Agrupado en"]))
        return creado

    # -- Movimientos
    def movimiento(self, fila):
        tipo = _opcion(fila["Tipo"], TIPOS_MOV, "El tipo")
        fecha = _fecha(fila["Fecha"])
        if fecha > self.hoy:
            raise ValueError("La fecha está en el futuro.")
        cuenta = self._buscar_cuenta(fila["Cuenta"])
        monto = _decimal(fila["Monto"], "el monto")
        concepto = None
        if _texto(fila["Concepto"]):
            concepto, creado = self._concepto(tipo, _texto(fila["Concepto"]).split("›")[-1].strip())
            if creado:
                self._sumar("Conceptos")
        miembro = None
        if _texto(fila["Miembro"]):
            miembro = self._buscar_miembro(fila["Miembro"])
            if miembro is None:
                raise ValueError(f"No existe el miembro «{_texto(fila['Miembro'])}». Agrégalo en la hoja Miembros.")
        datos = dict(
            tipo=tipo, fecha=fecha, cuenta=cuenta, monto=monto, concepto=concepto,
            descripcion=_texto(fila["Descripción"])[:200], miembro=miembro,
            tercero=_texto(fila["Tercero"])[:120], referencia=_texto(fila["Referencia"])[:60],
        )
        # Huella: mismo contenido + cuántas veces va repetido en el archivo.
        base = "|".join([str(self.o.pk), tipo, fecha.isoformat(), str(cuenta.pk), f"{monto.normalize():f}",
                         _clave(concepto.nombre if concepto else ""), _clave(datos["descripcion"]),
                         str(miembro.pk if miembro else ""), _clave(datos["tercero"]), _clave(datos["referencia"])])
        self._vistos[base] = self._vistos.get(base, 0) + 1
        huella = uuid.uuid5(ESPACIO, f"{base}|{self._vistos[base]}")
        movimiento = Movimiento(organizacion=self.o, uuid_cliente=huella, **datos)
        _, creado = registrar_movimiento(movimiento, membresia=self.membresia)
        return creado

    # -- Inventario
    def articulo(self, fila):
        nombre = _texto(fila["Artículo"])
        if not nombre:
            raise ValueError("Falta el nombre del artículo.")
        if Articulo.objects.filter(organizacion=self.o, nombre__iexact=nombre).exists():
            return False
        clase = _opcion(fila["Clase"], CLASES, "La clase")
        categoria = None
        if _texto(fila["Categoría"]):
            categoria, _ = CategoriaArticulo.objects.get_or_create(
                organizacion=self.o, nombre__iexact=_texto(fila["Categoría"]),
                defaults={"nombre": _texto(fila["Categoría"])[:80]})
        articulo = Articulo(
            organizacion=self.o, clase=clase, nombre=nombre, codigo=_texto(fila["Código"])[:40], categoria=categoria,
            unidad=_texto(fila["Unidad"])[:20] or "unidad",
            minimo=_decimal(fila["Mínimo"], "el mínimo", obligatorio=False) if clase == "consumible" else None,
            valor_unitario=_decimal(fila["Valor unitario"], "el valor unitario", obligatorio=False),
            moneda=_opcion(fila["Moneda"] or "USD", MONEDAS, "La moneda"),
        )
        articulo.full_clean()
        articulo.save()
        cantidad = _decimal(fila["Cantidad"], "la cantidad", obligatorio=False)
        if cantidad:
            lugar = _texto(fila["Ubicación"])
            if not lugar:
                raise ValueError("Con cantidad, la ubicación es obligatoria.")
            ubicacion = Ubicacion.objects.filter(organizacion=self.o, nombre__iexact=lugar).first()
            if ubicacion is None:
                ubicacion = Ubicacion.objects.create(organizacion=self.o, nombre=lugar[:80])
                self._sumar("Ubicaciones")
            inventario.registrar(MovimientoInventario(
                organizacion=self.o, articulo=articulo, tipo="entrada", motivo="inicial", fecha=self.hoy,
                cantidad=cantidad, destino=ubicacion, nota="Migración desde Excel",
            ), usuario=self.membresia.user)
        return True

    def correr(self, libro):
        self._vistos = {}
        self._cada(libro, "Miembros", self.miembro)
        self._cada(libro, "Cuentas", self.cuenta)
        self._cada(libro, "Conceptos", self.concepto)
        self._cada(libro, "Movimientos", self.movimiento)
        self._cada(libro, "Inventario", self.articulo)
        return self.r


def importar(organizacion, contenido, *, membresia, simular=True):
    """Lee la plantilla. Devuelve un `Resultado`; `guardado` es True solo si
    no era simulación y no hubo ningún error."""
    if len(contenido) > MAX_BYTES:
        raise ArchivoInvalido("El archivo pesa más de 2 MB. Divídelo en varios.")
    try:
        libro = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - cualquier cosa que no sea un .xlsx legible
        raise ArchivoInvalido("No se pudo leer el archivo. Debe ser un Excel (.xlsx) hecho con la plantilla.")
    if not any(h in libro.sheetnames for h in HOJAS):
        raise ArchivoInvalido("El archivo no tiene ninguna de las hojas de la plantilla.")
    with transaction.atomic():
        resultado = _Importador(organizacion, membresia).correr(libro)
        if simular or resultado.errores:
            transaction.set_rollback(True)
        else:
            resultado.guardado = True
    if not resultado.total and not resultado.errores:
        resultado.avisos.append("No hay nada nuevo que cargar: todo lo del archivo ya existe.")
    return resultado
