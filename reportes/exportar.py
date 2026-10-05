"""Salidas de un `Reporte`: Excel (openpyxl) y PDF (ReportLab), más el
comprobante de un movimiento. Colores y jerarquía iguales a la pantalla:
carbón, latón y marfil; verde y terracota solo para variaciones."""

from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from cambio.services import formatear

from .estructura import ENTERO, FECHA, GRUPO, MONTO, PORCENTAJE, SECCION, SUBTOTAL, TOTAL

CARBON, GRAFITO, PIEDRA = "1C1B1A", "3D3A37", "6F6A63"
LATON, MARFIL, LINO = "B08D4C", "F6F3EC", "E4DED2"
VERDE, TERRACOTA = "2F6B4F", "A4442F"
TONO = {"bien": VERDE, "mal": TERRACOTA}


def _texto_celda(valor, columna):
    """El valor ya formateado, para el PDF (y como respaldo en pantalla)."""
    if valor is None or valor == "":
        return ""
    if columna.tipo == MONTO:
        return formatear(valor)
    if columna.tipo == PORCENTAJE:
        return f"{formatear(valor, 1)} %"
    if columna.tipo == FECHA and isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    return str(valor)


# --- Excel -----------------------------------------------------------------


def a_excel(reporte, organizacion, generado_por=""):
    libro = Workbook()
    hoja = libro.active
    hoja.title = reporte.titulo[:31].replace("/", "-").replace(":", " ")
    n = len(reporte.columnas)
    ultima = get_column_letter(n)

    hoja["A1"] = organizacion.nombre
    hoja["A1"].font = Font(name="Calibri", size=14, bold=True, color=CARBON)
    hoja["A2"] = reporte.titulo
    hoja["A2"].font = Font(name="Calibri", size=12, bold=True, color=GRAFITO)
    hoja["A3"] = reporte.subtitulo
    hoja["A3"].font = Font(name="Calibri", size=10, color=PIEDRA)
    for fila in (1, 2, 3):
        hoja.merge_cells(f"A{fila}:{ultima}{fila}")

    fila_enc = 5
    borde_fino = Side(style="thin", color=LINO)
    for i, col in enumerate(reporte.columnas, start=1):
        celda = hoja.cell(row=fila_enc, column=i, value=col.titulo)
        celda.font = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
        celda.fill = PatternFill("solid", fgColor=CARBON)
        celda.alignment = Alignment(horizontal="right" if col.numerica else "left", vertical="center", wrap_text=True)
        hoja.column_dimensions[get_column_letter(i)].width = max(9, round(col.ancho * 11))
    hoja.row_dimensions[fila_enc].height = 22
    hoja.freeze_panes = hoja.cell(row=fila_enc + 1, column=1)

    r = fila_enc + 1
    for fila in reporte.filas:
        negrita = fila.clase in (GRUPO, SUBTOTAL, TOTAL, SECCION)
        for i, (valor, col) in enumerate(zip(fila.celdas, reporte.columnas), start=1):
            celda = hoja.cell(row=r, column=i)
            if isinstance(valor, Decimal):
                celda.value = float(valor)  # Excel no tiene decimales exactos; se muestra a 2
            else:
                celda.value = valor
            if col.tipo == MONTO:
                celda.number_format = "#,##0.00;[Red]-#,##0.00"
            elif col.tipo == PORCENTAJE:
                celda.number_format = '0.0" %"'
            elif col.tipo == ENTERO:
                celda.number_format = "#,##0"
            elif col.tipo == FECHA:
                celda.number_format = "DD/MM/YYYY"
            celda.alignment = Alignment(
                horizontal="right" if col.numerica else "left", vertical="center",
                indent=fila.sangria if i == 1 else 0,
            )
            color = TONO.get(fila.tonos.get(i - 1), CARBON)
            celda.font = Font(name="Calibri", size=10, bold=negrita, color="FFFFFF" if fila.clase == SECCION else color)
            if fila.clase == SECCION:
                celda.fill = PatternFill("solid", fgColor=GRAFITO)
            elif fila.clase == TOTAL:
                celda.fill = PatternFill("solid", fgColor=MARFIL)
                celda.border = Border(top=Side(style="thin", color=CARBON), bottom=Side(style="thin", color=CARBON))
            elif fila.clase == SUBTOTAL:
                celda.border = Border(top=borde_fino)
            else:
                celda.border = Border(bottom=borde_fino)
        r += 1

    r += 1
    for nota in reporte.notas:
        hoja.cell(row=r, column=1, value=nota).font = Font(name="Calibri", size=9, italic=True, color=PIEDRA)
        hoja.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n)
        r += 1
    pie = f"Generado por Arca el {timezone.localtime():%d/%m/%Y %H:%M}" + (f" · {generado_por}" if generado_por else "")
    hoja.cell(row=r + 1, column=1, value=pie).font = Font(name="Calibri", size=9, color=PIEDRA)

    hoja.page_setup.orientation = "landscape" if reporte.horizontal else "portrait"
    hoja.page_setup.fitToWidth, hoja.page_setup.fitToHeight = 1, 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.print_title_rows = f"{fila_enc}:{fila_enc}"

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


# --- PDF -------------------------------------------------------------------


def _hex(c):
    return colors.HexColor("#" + c)


ESTILO = {
    "org": ParagraphStyle("org", fontName="Helvetica-Bold", fontSize=13, leading=16, textColor=_hex(CARBON)),
    "titulo": ParagraphStyle("titulo", fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=_hex(GRAFITO)),
    "sub": ParagraphStyle("sub", fontName="Helvetica", fontSize=8.5, leading=11, textColor=_hex(PIEDRA)),
    "nota": ParagraphStyle("nota", fontName="Helvetica-Oblique", fontSize=7.5, leading=10, textColor=_hex(PIEDRA)),
    "celda": ParagraphStyle("celda", fontName="Helvetica", fontSize=8, leading=10, textColor=_hex(CARBON)),
    "firma": ParagraphStyle("firma", fontName="Helvetica", fontSize=8, leading=10, textColor=_hex(PIEDRA), alignment=TA_CENTER),
}


def _encabezado(organizacion, titulo, subtitulo, ancho):
    """Logo de Arca a la izquierda; organización, título y período a la derecha."""
    textos = [Paragraph(_esc(organizacion.nombre), ESTILO["org"])]
    datos = " · ".join(x for x in (f"RIF {organizacion.rif}" if organizacion.rif else "", organizacion.telefono) if x)
    if datos:
        textos.append(Paragraph(_esc(datos), ESTILO["sub"]))
    textos += [Spacer(1, 4), Paragraph(_esc(titulo), ESTILO["titulo"]), Paragraph(_esc(subtitulo), ESTILO["sub"])]
    logo = Path(settings.BASE_DIR) / "static" / "img" / "logo.png"
    if logo.exists():
        imagen = Image(str(logo), width=1.75 * cm, height=1.75 * cm * 435 / 480)
        tabla = Table([[imagen, textos]], colWidths=[2.3 * cm, ancho - 2.3 * cm])
    else:
        tabla = Table([[textos]], colWidths=[ancho])
    tabla.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("LINEBELOW", (0, 0), (-1, -1), 1.2, _hex(LATON)), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    return tabla


def _esc(texto):
    return str(texto).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _pie(generado_por):
    marca = f"Generado por Arca el {timezone.localtime():%d/%m/%Y %H:%M}" + (f" · {generado_por}" if generado_por else "")

    def dibujar(lienzo, doc):
        lienzo.saveState()
        lienzo.setFont("Helvetica", 7.5)
        lienzo.setFillColor(_hex(PIEDRA))
        lienzo.drawString(doc.leftMargin, 0.9 * cm, marca)
        lienzo.drawRightString(doc.pagesize[0] - doc.rightMargin, 0.9 * cm, f"Página {doc.page}")
        lienzo.restoreState()

    return dibujar


def _firmas(etiquetas, ancho):
    if not etiquetas:
        return []
    celda = ancho / len(etiquetas)
    tabla = Table([[Paragraph("_" * 30 + f"<br/>{_esc(e)}", ESTILO["firma"]) for e in etiquetas]], colWidths=[celda] * len(etiquetas))
    return [Spacer(1, 1.4 * cm), tabla]


def a_pdf(reporte, organizacion, generado_por=""):
    pagina = landscape(letter) if reporte.horizontal else letter
    margen = 1.5 * cm
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=pagina, leftMargin=margen, rightMargin=margen, topMargin=1.3 * cm,
                            bottomMargin=1.6 * cm, title=reporte.titulo, author="Arca")
    ancho = pagina[0] - 2 * margen
    total_anchos = sum(c.ancho for c in reporte.columnas)
    anchos = [ancho * c.ancho / total_anchos for c in reporte.columnas]

    def estilo(columna, negrita=False, color=CARBON):
        return ParagraphStyle(
            "c", parent=ESTILO["celda"], fontName="Helvetica-Bold" if negrita else "Helvetica",
            alignment=TA_RIGHT if columna.numerica else 0, textColor=_hex(color),
        )

    datos = [[Paragraph(_esc(c.titulo), estilo(c, True, "FFFFFF")) for c in reporte.columnas]]
    reglas = [
        ("BACKGROUND", (0, 0), (-1, 0), _hex(CARBON)), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("LINEBELOW", (0, 1), (-1, -1), 0.3, _hex(LINO)),
    ]
    for n_fila, fila in enumerate(reporte.filas, start=1):
        negrita = fila.clase in (GRUPO, SUBTOTAL, TOTAL, SECCION)
        celdas = []
        for i, (valor, col) in enumerate(zip(fila.celdas, reporte.columnas)):
            texto = _esc(_texto_celda(valor, col))
            if i == 0 and fila.sangria:
                texto = "&nbsp;" * (3 * fila.sangria) + texto
            color = "FFFFFF" if fila.clase == SECCION else TONO.get(fila.tonos.get(i), CARBON)
            celdas.append(Paragraph(texto, estilo(col, negrita, color)))
        datos.append(celdas)
        if fila.clase == SECCION:
            reglas.append(("BACKGROUND", (0, n_fila), (-1, n_fila), _hex(GRAFITO)))
        elif fila.clase == TOTAL:
            reglas += [("BACKGROUND", (0, n_fila), (-1, n_fila), _hex(MARFIL)),
                       ("LINEABOVE", (0, n_fila), (-1, n_fila), 0.7, _hex(CARBON))]
        elif fila.clase == SUBTOTAL:
            reglas.append(("LINEABOVE", (0, n_fila), (-1, n_fila), 0.4, _hex(PIEDRA)))

    tabla = Table(datos, colWidths=anchos, repeatRows=1)
    tabla.setStyle(TableStyle(reglas))
    partes = [_encabezado(organizacion, reporte.titulo, reporte.subtitulo, ancho), Spacer(1, 10), tabla, Spacer(1, 8)]
    partes += [Paragraph(_esc(nota), ESTILO["nota"]) for nota in reporte.notas]
    partes += _firmas(reporte.firmas, ancho)
    pie = _pie(generado_por)
    doc.build(partes, onFirstPage=pie, onLaterPages=pie)
    return buffer.getvalue()


# --- Comprobante de un movimiento ------------------------------------------


def comprobante_pdf(movimiento, organizacion, generado_por=""):
    """Comprobante de ingreso o de egreso: media carta de contenido, con firmas."""
    m = movimiento
    margen = 2 * cm
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=margen, rightMargin=margen, topMargin=1.5 * cm,
                            bottomMargin=1.6 * cm, title=f"Comprobante {m.numero_vale}", author="Arca")
    ancho = letter[0] - 2 * margen
    simbolo = {"USD": "$", "VES": "Bs."}[m.moneda]
    otra = f"Bs. {formatear(m.monto_ves)}" if m.moneda == "USD" else f"$ {formatear(m.monto_usd)}"

    etiqueta = ParagraphStyle("e", parent=ESTILO["sub"], fontSize=8)
    valor = ParagraphStyle("v", parent=ESTILO["celda"], fontSize=10, leading=13)
    grande = ParagraphStyle("g", parent=ESTILO["org"], fontSize=20, leading=24, alignment=TA_RIGHT)

    def par(nombre, contenido):
        return [Paragraph(nombre.upper(), etiqueta), Paragraph(_esc(contenido or "—"), valor)]

    tipo = "INGRESO" if m.es_ingreso else "EGRESO"
    cabecera = Table([[
        [Paragraph(f"COMPROBANTE DE {tipo}", ESTILO["titulo"]), Paragraph(f"N.º {m.numero_vale:05d} · {m.caja.nombre} · Ejercicio {m.ejercicio.nombre}", ESTILO["sub"])],
        [Paragraph(f"{simbolo} {formatear(m.monto)}", grande), Paragraph(f"Equivale a {otra} · tasa Bs. {formatear(m.tasa_aplicada)}", ParagraphStyle("d", parent=ESTILO["sub"], alignment=TA_RIGHT))],
    ]], colWidths=[ancho * 0.55, ancho * 0.45])
    cabecera.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))

    detalle = Table([
        [par("Fecha", f"{m.fecha:%d/%m/%Y}"), par("Cuenta", f"{m.cuenta.nombre} ({m.moneda})")],
        [par("Concepto", m.titulo), par("Recibido de" if m.es_ingreso else "Entregado a", m.nombre_persona)],
        [par("Descripción", m.descripcion if m.concepto_id else ""), par("Referencia", m.referencia)],
        [par("Registrado por", m.registrado_por.nombre_para_mostrar()), par("Estado", m.get_estado_display())],
    ], colWidths=[ancho / 2] * 2)
    detalle.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LINEBELOW", (0, 0), (-1, -1), 0.3, _hex(LINO)),
    ]))

    partes = [_encabezado(organizacion, "Comprobante contable", "Documento interno de la organización", ancho),
              Spacer(1, 14), cabecera, Spacer(1, 10), detalle]
    if m.estado == m.Estado.ANULADO:
        partes += [Spacer(1, 10), Paragraph(
            f"<b>ANULADO</b> el {timezone.localtime(m.fecha_anulacion):%d/%m/%Y}. Motivo: {_esc(m.motivo_anulacion)}",
            ParagraphStyle("a", parent=ESTILO["celda"], textColor=_hex(TERRACOTA), fontSize=10, leading=13))]
    partes += _firmas(["Entregado por", "Recibido por", "Autorizado por"], ancho)
    pie = _pie(generado_por)
    doc.build(partes, onFirstPage=pie, onLaterPages=pie)
    return buffer.getvalue()
