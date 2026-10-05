from datetime import date
from decimal import Decimal as D
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import RegistroAuditoria
from finanzas.models import Traslado
from finanzas.services import registrar_traslado
from presupuestos.tests import PresupuestoMixin

from . import services
from .estructura import NORMAL, TOTAL

ENE, MAR = date(2026, 1, 1), date(2026, 3, 31)


class DatosMixin(PresupuestoMixin):
    """Enero–marzo de 2026 con cifras redondas, todo a tasa 100."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        self = cls()
        self.gasto(date(2026, 1, 10), "300", tipo="ingreso", miembro=cls.residente)
        self.gasto(date(2026, 2, 10), "100", tipo="ingreso", tercero="Donante")
        self.gasto(date(2026, 1, 15), "80")
        self.gasto(date(2026, 2, 15), "5000", cuenta=cls.ves)          # $50
        self.gasto(date(2026, 3, 15), "40", concepto=cls.luz, miembro=cls.residente)
        self.gasto(date(2026, 3, 16), "10", concepto=None, descripcion="Taxi")
        self.gasto(date(2026, 3, 17), "77").transicionar("anulado", cls.director.user, motivo="x")  # no cuenta
        registrar_traslado(Traslado(organizacion=cls.org, fecha=date(2026, 2, 20), cuenta_origen=cls.usd,
                                    cuenta_destino=cls.ves, monto_origen=D("20"), monto_destino=D("2100")),
                           usuario=cls.director.user)

    @staticmethod
    def fila(reporte, primera_celda):
        return next(f for f in reporte.filas if f.celdas[0] == primera_celda or (
            len(f.celdas) > 3 and primera_celda in (f.celdas[2], f.celdas[3])))


class ReportesTests(DatosMixin, TestCase):
    def test_estado_de_ingresos_y_egresos(self):
        r = services.estado_ingresos_egresos(self.org, ENE, MAR, None, "USD")
        self.assertEqual(self.fila(r, "Total ingresos").celdas[1], D("400"))
        self.assertEqual(self.fila(r, "Total egresos").celdas[1], D("180"))     # 80 + 50 + 40 + 10; el anulado no
        self.assertEqual(self.fila(r, "RESULTADO DEL PERÍODO (ingresos − egresos)").celdas[1], D("220"))
        self.assertEqual(self.fila(r, "Otros (sin concepto)").celdas[1], D("10"))
        comida = self.fila(r, "Comida")
        self.assertEqual(comida.celdas[1], D("130"))
        self.assertAlmostEqual(float(comida.celdas[2]), 72.22, places=2)
        en_bs = services.estado_ingresos_egresos(self.org, ENE, MAR, None, "VES")
        self.assertEqual(self.fila(en_bs, "Total egresos").celdas[1], D("18000"))

    def test_variacion_contra_el_periodo_anterior(self):
        r = services.comparativo_periodos(self.org, date(2026, 2, 1), date(2026, 2, 28), None, "USD")
        comida = self.fila(r, "Comida")
        self.assertEqual(comida.celdas[1:4], [D("50"), D("80"), D("-30")])
        self.assertEqual(comida.tonos[3], "bien")       # se gastó menos
        aportes = self.fila(r, "Aportes")
        self.assertEqual(aportes.celdas[1:4], [D("100"), D("300"), D("-200")])
        self.assertEqual(aportes.tonos[3], "mal")       # entró menos

    def test_por_miembro(self):
        r = services.por_miembro(self.org, ENE, MAR, None, "USD")
        residente = self.fila(r, "Residente")
        self.assertEqual(residente.celdas[2:6], [D("300"), D("40"), D("260"), 2])
        otros = self.fila(r, "Terceros y movimientos sin miembro")
        self.assertEqual(otros.celdas[2:4], [D("100"), D("140")])
        self.assertEqual(self.fila(r, "Total").celdas[2:5], [D("400"), D("180"), D("220")])

    def test_flujo_mensual(self):
        r = services.flujo_mensual(self.org, self.ejercicio, None, "USD")
        self.assertEqual(self.fila(r, "Enero 2026").celdas[1:5], [D("300"), D("80"), D("220"), D("220")])
        self.assertEqual(self.fila(r, "Marzo 2026").celdas[1:5], [D("0"), D("50"), D("-50"), D("220")])
        self.assertEqual(self.fila(r, "Total 2026").celdas[1:4], [D("400"), D("180"), D("220")])
        self.assertEqual(len([f for f in r.filas if f.clase == NORMAL]), 12)

    def test_libro_de_cuenta_con_saldo_corrido_y_traslados(self):
        r = services.libro_cuenta(self.usd, date(2026, 2, 1), MAR)
        self.assertEqual(r.filas[0].celdas[6], D("270"))             # 50 inicial + 300 − 80 en enero
        final = r.filas[-1]
        self.assertEqual((final.clase, final.celdas[4], final.celdas[5], final.celdas[6]), (TOTAL, D("100"), D("70"), D("300")))
        self.assertTrue(any("Traslado a" in str(f.celdas[2]) for f in r.filas))
        banco = services.libro_cuenta(self.ves, ENE, MAR)
        self.assertEqual(banco.filas[-1].celdas[6], D("-2900"))      # 0 − 5000 + 2100, en bolívares
        self.assertEqual(banco.filas[-1].tonos, {})

    def test_libro_diario(self):
        r = services.libro_diario(self.org, ENE, MAR, None, "USD")
        self.assertEqual(len([f for f in r.filas if f.clase == NORMAL]), 6)
        self.assertEqual(self.fila(r, "Totales").celdas[7:9], [D("400"), D("180")])

    def test_presupuesto_contra_real(self):
        r = services.presupuesto_vs_real(self.p, date(2026, 3, 1))
        comida = self.fila(r, "Comida")
        self.assertEqual(comida.celdas[1:4], [D("100"), D("0"), D("-100")])
        self.assertEqual(comida.celdas[5:8], [D("300"), D("130"), D("-170")])
        self.assertEqual(comida.tonos, {3: "bien", 7: "bien"})
        self.assertEqual(self.fila(r, "Luz (sin partida)").celdas[2], D("40"))
        self.assertTrue(r.horizontal)
        self.assertTrue(any("BORRADOR" in n for n in r.notas))

    def test_analisis(self):
        a = services.analisis(self.org, self.ejercicio, ENE, MAR, None, "USD")
        self.assertEqual((a["ingresos"], a["egresos"], a["resultado"]), (D("400"), D("180"), D("220")))
        self.assertEqual(a["top_egresos"][0]["nombre"], "Comida")
        self.assertEqual(a["top_egresos"][0]["ancho"], "100.0")
        self.assertEqual(a["aportantes"][0]["monto"], D("300"))
        self.assertEqual(a["mayor_egreso_monto"], D("80"))
        self.assertEqual(len(a["serie"]), 12)

    def test_filtro_por_caja(self):
        r = services.estado_ingresos_egresos(self.org, ENE, MAR, self.caja, "USD")
        self.assertEqual(self.fila(r, "Total egresos").celdas[1], D("180"))
        from finanzas.models import Caja
        vacia = Caja.objects.create(organizacion=self.org, nombre="Vacía")
        self.assertTrue(services.estado_ingresos_egresos(self.org, ENE, MAR, vacia, "USD").vacio)


class VistasReportesTests(DatosMixin, TestCase):
    Q = "?desde=2026-01-01&hasta=2026-03-31"

    def test_todas_las_pantallas_abren(self):
        self.entrar("dir_a")
        for url in [reverse("reportes:inicio"), reverse("reportes:analisis") + self.Q,
                    reverse("reportes:presupuesto", args=[self.p.pk]) + "?mes=2026-03",
                    *[reverse("reportes:reporte", args=[c]) + self.Q for c in ("estado", "variacion", "flujo", "miembros", "libro", "diario")]]:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 200)
        self.assertEqual(self.client.get(reverse("reportes:reporte", args=["inventado"])).status_code, 404)
        self.assertContains(self.client.get(reverse("reportes:analisis") + self.Q), "col-ingreso")

    def test_pdf_y_excel_de_cada_reporte(self):
        self.entrar("dir_a")
        urls = [reverse("reportes:reporte", args=[c]) + self.Q for c in ("estado", "variacion", "flujo", "miembros", "libro", "diario")]
        urls.append(reverse("reportes:presupuesto", args=[self.p.pk]) + "?mes=2026-03")
        for url in urls:
            with self.subTest(url=url):
                pdf = self.client.get(url + "&formato=pdf")
                self.assertEqual(pdf["Content-Type"], "application/pdf")
                self.assertTrue(pdf.content.startswith(b"%PDF"))
                xlsx = self.client.get(url + "&formato=xlsx")
                self.assertTrue(load_workbook(BytesIO(xlsx.content)).active.max_row > 5)
        self.assertEqual(RegistroAuditoria.objects.filter(accion="emitir_reporte", organizacion=self.org).count(), 14)
        self.assertEqual(self.client.get(urls[0] + "&formato=docx").status_code, 404)

    def test_el_excel_lleva_numeros_no_texto(self):
        self.entrar("dir_a")
        r = self.client.get(reverse("reportes:reporte", args=["estado"]) + self.Q + "&formato=xlsx")
        hoja = load_workbook(BytesIO(r.content)).active
        self.assertEqual(hoja["A1"].value, "Casa A")
        valores = {fila[0].value: fila[1].value for fila in hoja.iter_rows(min_row=6) if fila[0].value}
        self.assertEqual(valores["Total ingresos"], 400)
        self.assertEqual(valores["Total egresos"], 180)

    def test_comprobante(self):
        from finanzas.models import Movimiento
        m = Movimiento.objects.filter(organizacion=self.org).first()
        self.entrar("dir_a")
        r = self.client.get(reverse("reportes:comprobante", args=[m.pk]))
        self.assertTrue(r.content.startswith(b"%PDF"))
        self.entrar("comprador")   # solo registra: no ve lo que no registró él
        self.assertEqual(self.client.get(reverse("reportes:comprobante", args=[m.pk])).status_code, 404)
        self.entrar("residente")
        self.assertEqual(self.client.get(reverse("reportes:comprobante", args=[m.pk])).status_code, 403)

    def test_permisos(self):
        for usuario in ("comprador", "residente"):
            self.entrar(usuario)
            for url in (reverse("reportes:inicio"), reverse("reportes:analisis"), reverse("reportes:reporte", args=["estado"]),
                        reverse("reportes:presupuesto", args=[self.p.pk])):
                with self.subTest(usuario=usuario, url=url):
                    self.assertEqual(self.client.get(url).status_code, 403)
        # Ver el libro ya no basta: los reportes tienen su propia casilla en el tipo de miembro.
        self.entrar("tesorero")
        self.assertEqual(self.client.get(reverse("reportes:reporte", args=["estado"])).status_code, 403)
        self.assertEqual(self.client.get(reverse("finanzas:movimiento_lista")).status_code, 200)
        self.tesorero.tipo.puede_ver_reportes = True
        self.tesorero.tipo.save()
        self.assertEqual(self.client.get(reverse("reportes:reporte", args=["estado"])).status_code, 200)
        self.assertEqual(self.client.get(reverse("reportes:analisis")).status_code, 200)

    def test_aislamiento(self):
        self.entrar("dir_b")
        r = self.client.get(reverse("reportes:reporte", args=["estado"]) + self.Q)
        self.assertTrue(r.context["reporte"].vacio)
        self.assertNotContains(r, "Comida")
        a = self.client.get(reverse("reportes:analisis") + self.Q).context["a"]
        self.assertEqual((a["ingresos"], a["egresos"]), (D("0"), D("0")))
        # Pedir el libro de una cuenta ajena no la muestra: cae en una propia.
        r = self.client.get(reverse("reportes:reporte", args=["libro"]) + self.Q + f"&cuenta={self.usd.pk}")
        self.assertEqual(r.context["v"]["cuenta"], self.cuenta_b)
        r = self.client.get(reverse("reportes:reporte", args=["estado"]) + self.Q + f"&caja={self.caja.pk}")
        self.assertIsNone(r.context["v"]["caja"])
